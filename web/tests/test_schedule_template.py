"""The blank Excel template an ansvarig fills in instead of the web form."""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from duty_web.schedule_template import build_template
from duty_web.schedule_xlsx import read_schedule

PLAYERS = ["Alva Exempel", "Tova Exempel"]
ACTIVITIES = ["Bästkustcupen", "Cafépass"]


def _open(data: bytes):
    return load_workbook(BytesIO(data))


def _filled(data: bytes, rows) -> bytes:
    workbook = _open(data)
    sheet = workbook.worksheets[0]
    for number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(number, column, value)
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_the_schedule_sheet_comes_first_with_the_imports_headings():
    sheet = _open(build_template(PLAYERS, ACTIVITIES)).worksheets[0]

    assert sheet.title == "Schema"
    assert [c.value for c in sheet[1]] == [
        "Aktivitet", "Datum", "Tid", "Person 1", "Person 2", "Person 3",
    ]
    assert sheet.freeze_panes == "A2"


def test_a_filled_in_template_imports_as_written():
    data = _filled(build_template(PLAYERS, ACTIVITIES), [
        ["Cafépass", date(2026, 11, 14), "9-12", "Tova Exempel", "Alva Exempel"],
        ["Bästkustcupen", None, "Fre 18-21", "Alva Exempel"],
    ])

    rows = read_schedule(data)

    assert [(r["activity"], r["date"], r["start"], r["end"], r["player"]) for r in rows] == [
        ("Cafépass", "2026-11-14", "09:00", "12:00", "Tova Exempel"),
        ("Cafépass", "2026-11-14", "09:00", "12:00", "Alva Exempel"),
        ("Bästkustcupen", None, "18:00", "21:00", "Alva Exempel"),
    ]


def test_an_untouched_template_imports_nothing():
    """Neither the example on the guide sheet nor the dropdown lists are passes."""
    assert read_schedule(build_template(PLAYERS, ACTIVITIES)) == []


def test_times_are_text_so_excel_never_turns_9_12_into_a_date():
    sheet = _open(build_template(PLAYERS, ACTIVITIES)).worksheets[0]

    assert sheet["C2"].number_format == "@"
    assert sheet["B2"].number_format == "yyyy-mm-dd"


def _validations(sheet):
    return {str(v.sqref): v for v in sheet.data_validations.dataValidation}


def test_people_and_activities_are_offered_from_lists_that_allow_anything_else():
    workbook = _open(build_template(PLAYERS, ACTIVITIES))
    validations = _validations(workbook["Schema"])
    lists = workbook["Listor"]

    person = next(v for ref, v in validations.items() if ref.startswith("D2"))
    activity = next(v for ref, v in validations.items() if ref.startswith("A2"))
    assert person.type == activity.type == "list"
    assert [lists.cell(row, 1).value for row in range(1, 3)] == PLAYERS
    assert [lists.cell(row, 2).value for row in range(1, 3)] == ACTIVITIES
    assert "Listor!$A$1:$A$2" in person.formula1
    assert "Listor!$B$1:$B$2" in activity.formula1
    # A child or activity not in the list is Granska's question, not Excel's refusal.
    assert not person.showErrorMessage and not activity.showErrorMessage
    assert next(ref for ref in validations if ref.startswith("D2")) == "D2:F501"
    assert lists.sheet_state == "hidden"


def test_the_date_column_warns_about_anything_but_a_date():
    validations = _validations(_open(build_template(PLAYERS, ACTIVITIES))["Schema"])

    when = next(v for ref, v in validations.items() if ref.startswith("B2"))
    assert when.type == "date" and when.errorStyle == "warning" and when.allow_blank


def test_an_empty_list_offers_no_dropdown():
    validations = _validations(_open(build_template([], ACTIVITIES))["Schema"])

    assert not any(ref.startswith("D2") for ref in validations)


def test_the_guide_sheet_explains_the_columns_with_an_example():
    guide = _open(build_template(PLAYERS, ACTIVITIES))["Så fyller du i"]
    text = " ".join(str(c.value) for row in guide.iter_rows() for c in row if c.value)

    assert "Fre 18-21" in text and "tre" in text
    assert ACTIVITIES[0] in text  # the example uses an activity the team knows


def test_a_name_that_looks_like_a_formula_stays_text():
    """Names come from 360Player; Excel must never run one."""
    workbook = _open(build_template(["=HYPERLINK(\"http://x\")"], ["+cmd"]))

    assert workbook["Listor"]["A1"].data_type == "s"
    assert workbook["Listor"]["B1"].data_type == "s"
