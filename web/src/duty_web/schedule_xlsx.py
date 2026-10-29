"""The coach's season workbook, read in memory (spec: Excel import).

Replaces the old scripts/build_seed_xlsx.py, so an ansvarig can import the
coach's own file. Two layouts are understood:

- the coach's: a sheet whose header row has Aktivitet, Datum, Tid and
  Person 1..3, one row per shift;
- a seed-format file: syssla, datum, tid and
  barn (with arena, station and anteckning), one row per person.

Either way the result is one row per seat, as plain values. The bytes are
read with openpyxl from memory and never written to disk.
"""

from __future__ import annotations

import itertools
import re
import warnings
import zipfile
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from .duties import canonical_duty_name

MAX_BYTES = 1_000_000
MAX_ROWS = 3000
# An .xlsx is a zip: these bound what it unpacks to, which MAX_BYTES doesn't.
MAX_MEMBERS = 100
MAX_UNZIPPED = 20_000_000
MAX_SHEETS = 10
MAX_COLUMNS = 30
MAX_TEXT = 200
UNREADABLE = "Filen går inte att läsa som Excel (.xlsx)."
PLACEHOLDER = re.compile(r"datum\s+ej\s+satt\b", re.IGNORECASE)
PERSON_COLUMNS = ("Person 1", "Person 2", "Person 3")
WEEKDAY_ABBREVIATIONS = {
    "mån": "Måndagar", "tis": "Tisdagar", "ons": "Onsdagar", "tor": "Torsdagar",
    "fre": "Fredagar", "lör": "Lördagar", "sön": "Söndagar",
}


class ScheduleFileError(ValueError):
    """The workbook can't be read. The message is shown next to the upload."""


def normalize_time_range(raw: object) -> tuple[str, str]:
    """'Fre 18-21' / '11:30 -16:00' / '9-12' → ('18:00', '21:00') etc.

    The sheet is hand-maintained, so the same shift shape appears with a
    weekday prefix, hour-only endpoints, dots, and any spacing around the dash.
    """
    text = " ".join(str(raw or "").split())
    text = re.sub(r"^[^0-9]+", "", text)
    parts = re.split(r"\s*[-–]\s*", text)
    if len(parts) != 2:
        raise ValueError(f"ogiltigt tidsintervall: {raw!r}")
    out = []
    for part in parts:
        hour, _, minute = part.strip().replace(".", ":").partition(":")
        if not hour.isdigit() or (minute and not minute.isdigit()):
            raise ValueError(f"ogiltigt tidsintervall: {raw!r}")
        h, m = int(hour), int(minute or 0)
        if h > 23 or m > 59:
            raise ValueError(f"ogiltigt tidsintervall: {raw!r}")
        out.append(f"{h:02d}:{m:02d}")
    return out[0], out[1]


def _weekday_prefix(raw: object) -> str:
    """'Fre 18-21' → 'Fredagar'; a plain time range → ''."""
    match = re.match(r"\s*([A-Za-zÅÄÖåäö]+)", str(raw or ""))
    if match is None:
        return ""
    return WEEKDAY_ABBREVIATIONS.get(match.group(1)[:3].casefold(), "")


def _date_cell(raw: object, number: int) -> date | None:
    """A date cell or ISO text → that date; blank or 'Datum ej satt (VT27)' → None.

    Anything else is refused with its row: a typo silently read as "no date
    yet" would replace a correctly dated shift and drop its swaps.
    """
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = _text(raw)
    if not text or PLACEHOLDER.match(text):
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        raise ScheduleFileError(f"Rad {number}: ogiltigt datum '{text[:40]}'") from None


def _text(raw: object) -> str:
    return " ".join(str(raw).split()) if raw is not None else ""


def _check_archive(data: bytes) -> None:
    """Refuse a workbook that unpacks to far more than it weighs.

    MAX_BYTES bounds the compressed upload; a deflated member can expand a
    thousandfold. zipfile never inflates a member past the size the archive
    declares for it (a lie ends in a CRC error), so checking the declared
    sizes bounds what openpyxl can be made to read.
    """
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        raise ScheduleFileError(UNREADABLE) from None
    if len(members) > MAX_MEMBERS or sum(m.file_size for m in members) > MAX_UNZIPPED:
        raise ScheduleFileError(
            "Filen är för stor uppackad. Spara en kopia med bara schemat och försök igen."
        )


def _times(raw: object, number: int) -> tuple[str, str]:
    try:
        return normalize_time_range(raw)
    except ValueError:
        raise ScheduleFileError(f"Rad {number}: ogiltig tid '{_text(raw)}'") from None


def _row(number, day, times, activity, player, venue="", station="", note="") -> dict:
    if any(len(text) > MAX_TEXT for text in (activity, player, venue, station, note)):
        raise ScheduleFileError(f"Rad {number}: för lång text (högst {MAX_TEXT} tecken per cell)")
    return {
        "row": number, "date": day.isoformat() if day else None,
        "start": times[0], "end": times[1], "activity": canonical_duty_name(activity),
        "player": player, "venue": venue, "station": station, "note": note,
    }


def _cell(row: tuple, index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def _coach_rows(rows: list[tuple]) -> list[dict] | None:
    header_at = next(
        (i for i, row in enumerate(rows) if {"Aktivitet", "Datum", "Tid"} <= {_text(c) for c in row}),
        None,
    )
    if header_at is None:
        return None
    columns = {_text(c): i for i, c in enumerate(rows[header_at]) if _text(c)}
    people = [columns[name] for name in PERSON_COLUMNS if name in columns]
    if not people:
        return None
    out = []
    for number, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        activity = _text(_cell(row, columns["Aktivitet"]))
        if not activity:
            continue
        raw_time = _cell(row, columns["Tid"])
        day = _date_cell(_cell(row, columns["Datum"]), number)
        times = _times(raw_time, number)
        # An undated shift still says which weekday it lands on, the one
        # thing a parent can plan around until the term's dates are set.
        note = "" if day else _weekday_prefix(raw_time)
        for index in people:
            player = _text(_cell(row, index))
            if player:
                out.append(_row(number, day, times, activity, player, note=note))
    return out


def _seed_rows(rows: list[tuple]) -> list[dict] | None:
    if not rows:
        return None
    columns = {_text(c).casefold(): i for i, c in enumerate(rows[0]) if _text(c)}
    if not {"syssla", "tid", "barn"} <= columns.keys():
        return None
    out = []
    for number, row in enumerate(rows[1:], start=2):
        if not any(cell is not None and _text(cell) for cell in row):
            continue
        activity = _text(_cell(row, columns["syssla"]))
        player = _text(_cell(row, columns["barn"]))
        if not activity or not player:
            raise ScheduleFileError(f"Rad {number}: saknar syssla eller barn")
        out.append(_row(
            number, _date_cell(_cell(row, columns.get("datum")), number),
            _times(_cell(row, columns["tid"]), number), activity, player,
            venue=_text(_cell(row, columns.get("arena"))),
            station=_text(_cell(row, columns.get("station"))),
            note=_text(_cell(row, columns.get("anteckning"))),
        ))
    return out


def read_schedule(data: bytes) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise ScheduleFileError("Filen är för stor (högst 1 MB).")
    _check_archive(data)
    # The file is untrusted, so the exception types a damaged archive can raise
    # (zlib, EOF, XML, openpyxl's own) are not enumerated: any failure while
    # reading it means "unreadable". Warnings from crafted files are silenced
    # so they neither reach the logs nor the test output.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except ScheduleFileError:
            raise
        except Exception:
            raise ScheduleFileError(UNREADABLE) from None
        try:
            return _read_sheets(workbook)
        except ScheduleFileError:
            raise
        except Exception:
            # A sheet is parsed as it's iterated, so the damage can surface here.
            raise ScheduleFileError(UNREADABLE) from None
        finally:
            workbook.close()


def _read_sheets(workbook) -> list[dict]:
    for sheet in workbook.worksheets[:MAX_SHEETS]:
        # Read one row past the cap, so a file that's too long is refused
        # rather than silently cut short. max_col keeps a cell far to the
        # right from widening every row.
        rows = list(itertools.islice(
            sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True), MAX_ROWS + 101
        ))
        if len(rows) > MAX_ROWS + 100:
            raise ScheduleFileError(f"Högst {MAX_ROWS} pass per fil.")
        found = _coach_rows(rows)
        if found is None:
            found = _seed_rows(rows)
        if found is not None:
            if len(found) > MAX_ROWS:
                raise ScheduleFileError(f"Högst {MAX_ROWS} pass per fil.")
            return found
    raise ScheduleFileError(
        "Hittade inget schema i filen: varken kolumnerna Aktivitet, Datum, Tid och "
        "Person 1 eller syssla, tid och barn."
    )
