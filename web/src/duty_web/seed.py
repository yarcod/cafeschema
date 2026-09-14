"""Schedule xlsx -> database import. Replaces every slot, keeps the roster.

Slots are derived data: the trainer's spreadsheet is their only source, so
re-importing it wholesale is how a corrected schedule gets in. Players,
parents and the links between them are *not* touched here — they come from
the roster import (roster.py), which is authoritative and additive, so a
re-import never disturbs a logged-in parent or a pending swap's owners.

A slot names its player in the 'barn' column; who may act on it follows
from that player's parents.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .models import Player, Slot, Team
from .roster import normalize_name


class SeedError(Exception):
    """Raised when the xlsx cannot be imported. Message is operator-facing."""


def _parse_time_range(raw: object, *, row_number: int) -> tuple[str, str]:
    parts = str(raw).split("-")
    if len(parts) != 2:
        raise SeedError(f"Rad {row_number}: ogiltigt tidsintervall '{raw}'")
    start, end = (p.strip() for p in parts)
    return start, end


def import_schedule(session: Session, path: str | Path, *, team_id: int) -> int:
    path = Path(path)
    if not path.is_file():
        raise SeedError(f"Schemafilen hittades inte: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise SeedError(f"Schemafilen är tom: {path}") from None

        columns = [str(c).strip().lower() if c is not None else "" for c in header]
        if not any(columns):
            raise SeedError(f"Schemafilen är tom: {path}")

        records = [
            {col: value for col, value in zip(columns, row) if col}
            for row in rows
            if any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()

    # Nothing else creates a Team row, and SQLite's default connection
    # settings don't enforce foreign keys, so importing against a team_id
    # that doesn't exist would otherwise silently write slots with a
    # dangling team_id. Fail loudly, and before wiping the existing slots
    # below, rather than after.
    if session.get(Team, team_id) is None:
        raise SeedError(
            f"Ingen grupp med team_id={team_id} finns. Skapa gruppen innan import."
        )

    # Every slot's owner is resolved through this map, so an empty roster
    # would import the whole season unassigned and nobody could log in to
    # anything. That is always an operator mistake, not a real schedule.
    players_by_name = {
        normalize_name(player.name): player
        for player in session.query(Player).filter(Player.team_id == team_id).all()
    }
    if not players_by_name:
        raise SeedError(
            "Inga spelare finns i gruppen. Importera laglistan (roster) före schemat."
        )

    session.query(Slot).delete()
    session.flush()

    count = 0
    try:
        for row_number, record in enumerate(records, start=2):
            player_name = " ".join(str(record.get("barn") or "").split())
            if not player_name:
                raise SeedError(f"Rad {row_number}: saknar barn")

            # A player who isn't on the roster yet still has a real shift —
            # it's just one nobody can act on until the roster catches up
            # (Slot.player_id is nullable for exactly this reason). Keep the
            # intended name visible via note rather than silently losing it.
            player = players_by_name.get(normalize_name(player_name))

            raw_date = record.get("datum")
            raw_date_str = "" if raw_date is None else str(raw_date).strip()
            if isinstance(raw_date, datetime):
                slot_date = raw_date.date()
            elif not raw_date_str:
                # Blank means "not yet scheduled" (e.g. a future term whose
                # exact dates aren't fixed yet) rather than a bad row.
                slot_date = None
            else:
                try:
                    slot_date = datetime.strptime(raw_date_str, "%Y-%m-%d").date()
                except ValueError:
                    raise SeedError(f"Rad {row_number}: ogiltigt datum '{raw_date}'")

            start_str, end_str = _parse_time_range(
                record.get("tid", ""), row_number=row_number
            )

            try:
                start_time = datetime.strptime(start_str, "%H:%M").time()
            except ValueError:
                raise SeedError(f"Rad {row_number}: ogiltigt starttid '{start_str}'")

            try:
                end_time = datetime.strptime(end_str, "%H:%M").time()
            except ValueError:
                raise SeedError(f"Rad {row_number}: ogiltigt sluttid '{end_str}'")

            anteckning = record.get("anteckning")
            note = (str(anteckning).strip() or None) if anteckning else None
            if player is None:
                unfilled_note = f"Ej matchad spelare för '{player_name}'"
                note = f"{unfilled_note} | {note}" if note else unfilled_note

            session.add(
                Slot(
                    team_id=team_id,
                    date=slot_date,
                    start_time=start_time,
                    end_time=end_time,
                    station=str(record.get("station") or "").strip(),
                    duty_name=str(record.get("syssla") or "").strip(),
                    venue=str(record.get("arena") or "").strip(),
                    child_name=player_name,
                    note=note,
                    player_id=player.id if player is not None else None,
                )
            )
            count += 1
    except SeedError:
        session.rollback()
        raise

    session.commit()
    return count
