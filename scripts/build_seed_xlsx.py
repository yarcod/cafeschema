"""Build a duty_web-seedable xlsx from a trainer's roster + parent contacts.

The trainer's schedule (`Säsong 26-27` sheet) lists one row per shift, with
up to three *players* staffing it, but duty_web needs one parent e-post per
slot so that a Person can log in. This script joins the two and flattens
each shift into one row per staffing player, using the policy agreed with
the team lead:

- Household with two parents on file -> first parent listed in the
  mailing-list CSV (arbitrary but simple; wrong-parent cases can swap or
  forward the login link once live).
- Player name not found under any spelling -> row is kept with no e-post;
  duty_web's importer leaves such a slot unassigned rather than dropping
  it (Slot.person_id is nullable for exactly this case).
- Shift with no date fixed yet ("Datum ej satt (VT27)") -> row is kept
  with a blank date; duty_web shows these under "Datum ej satt" instead of
  on the calendar, so a parent still sees what they're signed up for.

Spelningar som skiljer sig mellan schemat och kontaktlistan rättas via
data/name_fixes.json (se load_name_fixes) — den filen innehåller
personuppgifter och ligger utanför repot.

Usage:
    python scripts/build_seed_xlsx.py \\
        data/Bemanningsschema_26_27_Hans_260913_v2.xlsx \\
        data/f15_parent_mailing_list.csv \\
        data/seed_sasong_26_27.xlsx
"""

from __future__ import annotations

import csv
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

SHEET = "Säsong 26-27"
PERSON_COLUMNS = ("Person 1", "Person 2", "Person 3")

# Spelarnamn stavas olika i schemat och i kontaktlistan. Kopplingen mellan
# stavningarna är personuppgifter och ligger därför hos rosterfilerna, inte i
# koden: data/name_fixes.json, t.ex. {"Namn i schemat": "Namn i kontaktlistan"}.
NAME_FIXES_PATH = Path("data/name_fixes.json")


def load_name_fixes() -> dict[str, str]:
    if not NAME_FIXES_PATH.is_file():
        return {}
    return json.loads(NAME_FIXES_PATH.read_text(encoding="utf-8"))

# Only the arena-hosting duties have a venue in the source; the cup and
# market activities are held elsewhere and the sheet doesn't say where, so
# those are left blank rather than guessed at.
ARENA_VENUE = "Wallenstam arena"

SEED_HEADER = [
    "ar", "syssla", "arena", "station", "vecka", "datum", "veckodag",
    "tid", "namn", "epost", "barn", "anteckning",
]

WEEKDAYS_SV = (
    "Måndag", "Tisdag", "Onsdag", "Torsdag", "Fredag", "Lördag", "Söndag",
)


def _norm(name: str) -> str:
    return unicodedata.normalize("NFKC", " ".join(name.split())).casefold()


def _parse_date_cell(raw: object) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    try:
        return datetime.strptime(str(raw).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


WEEKDAY_ABBREVIATIONS = {
    "mån": "Måndagar", "tis": "Tisdagar", "ons": "Onsdagar", "tor": "Torsdagar",
    "fre": "Fredagar", "lör": "Lördagar", "sön": "Söndagar",
}


def _weekday_prefix(raw: object) -> str:
    """'Fre 18-21' -> 'Fredagar'; a plain time range -> ''."""
    match = re.match(r"\s*([A-Za-zÅÄÖåäö]+)", str(raw))
    if match is None:
        return ""
    return WEEKDAY_ABBREVIATIONS.get(match.group(1)[:3].casefold(), "")


def normalize_time_range(raw: object) -> str:
    """'Fre 18-21' / '11:30 -16:00' -> '18:00-21:00' / '11:30-16:00'.

    The sheet is hand-maintained, so the same shift shape appears with a
    weekday prefix, hour-only endpoints, and inconsistent spacing around
    the dash.
    """
    text = " ".join(str(raw).split())
    text = re.sub(r"^[^0-9]+", "", text)
    parts = re.split(r"\s*[-–]\s*", text)
    if len(parts) != 2:
        raise ValueError(f"ogiltigt tidsintervall: {raw!r}")
    normalized = []
    for part in parts:
        part = part.strip()
        if ":" not in part:
            part += ":00"
        hour, minute = part.split(":", 1)
        normalized.append(f"{int(hour):02d}:{int(minute):02d}")
    return "-".join(normalized)


def load_parents_by_player(csv_path: Path) -> dict[str, list[tuple[str, str]]]:
    """Player name (normalized) -> parents in file order (first = primary)."""
    parents: dict[str, list[tuple[str, str]]] = {}
    with csv_path.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            player = _norm(row["children_on_team"])
            parents.setdefault(player, []).append(
                (row["parent_name"].strip(), row["parent_email"].strip())
            )
    return parents


def build_rows(
    schedule_path: Path, parents_by_player: dict[str, list[tuple[str, str]]]
) -> list[list]:
    name_fixes = load_name_fixes()
    workbook = load_workbook(schedule_path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_row = next(
        row for row in rows if row and "Aktivitet" in [str(c).strip() if c else "" for c in row]
    )
    header = [str(c).strip() if c is not None else "" for c in header_row]
    idx = {name: i for i, name in enumerate(header)}

    out: list[list] = []
    unmatched: list[str] = []
    for row in rows[rows.index(header_row) + 1:]:
        activity = row[idx["Aktivitet"]]
        if not activity:
            continue
        activity = str(activity).strip()

        # Some rows are real date cells, others plain "YYYY-MM-DD" text, and
        # the rest a placeholder such as "Datum ej satt (VT27)" — the shift
        # is real and staffed either way, only the date is still open.
        slot_date = _parse_date_cell(row[idx["Datum"]])
        if slot_date is not None:
            datum = slot_date.isoformat()
            vecka = slot_date.isocalendar().week
            veckodag = WEEKDAYS_SV[slot_date.weekday()]
        else:
            datum, vecka, veckodag = "", "", ""

        tid = normalize_time_range(row[idx["Tid"]])
        # An undated shift still says which weekday it lands on ("Fre 18-21"),
        # which is the only thing a parent can plan around until the term's
        # dates are set — so keep it rather than dropping it with the prefix.
        shift_weekday = "" if slot_date is not None else _weekday_prefix(row[idx["Tid"]])
        arena = ARENA_VENUE if activity.startswith("Arena värdskap") else ""

        for column in PERSON_COLUMNS:
            raw_player = row[idx[column]]
            if not raw_player:
                continue
            player = " ".join(str(raw_player).split())
            player = name_fixes.get(player, player)

            matches = parents_by_player.get(_norm(player), [])
            if matches:
                namn, epost = matches[0]
            else:
                namn, epost = player, ""
                unmatched.append(player)
            note = shift_weekday

            out.append([
                slot_date.year if slot_date is not None else "",
                activity,
                arena,
                "",
                vecka,
                datum,
                veckodag,
                tid,
                namn,
                epost,
                player,
                note,
            ])

    if unmatched:
        print(
            f"Varning: {len(unmatched)} pass utan matchad kontakt, importeras otilldelade: "
            + ", ".join(sorted(set(unmatched))),
            file=sys.stderr,
        )

    return out


def write_seed_xlsx(rows: list[list], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(SEED_HEADER)
    for row in rows:
        ws.append(row)
    wb.save(out_path)


def main(argv: list[str]) -> int:
    if len(argv) != 4:
        print(__doc__)
        return 1
    schedule_path, contacts_path, out_path = (Path(a) for a in argv[1:])

    parents_by_player = load_parents_by_player(contacts_path)
    rows = build_rows(schedule_path, parents_by_player)
    write_seed_xlsx(rows, out_path)
    print(f"Skrev {len(rows)} pass till {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
