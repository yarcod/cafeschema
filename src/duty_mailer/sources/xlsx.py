"""Local .xlsx roster."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..models import Occurrence
from .rows import RosterError, occurrences_from_rows


class XlsxSource:
    """Reads the roster from the first worksheet of a local .xlsx file."""

    def __init__(self, path: str | Path, *, default_lead_days: tuple[int, ...]):
        self._path = Path(path)
        self._default_lead_days = default_lead_days

    def fetch(self) -> list[Occurrence]:
        if not self._path.is_file():
            raise RosterError(f"Schemafilen hittades inte: {self._path}")

        # read_only for large sheets; data_only so formula cells yield values.
        workbook = load_workbook(self._path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                raise RosterError(f"Schemafilen är tomt: {self._path}") from None

            columns = [str(c).strip().lower() if c is not None else "" for c in header]
            if not any(columns):
                raise RosterError(f"Schemafilen är tomt: {self._path}")

            dicts = [
                {col: value for col, value in zip(columns, row) if col}
                for row in rows
                if any(value is not None for value in row)
            ]
        finally:
            workbook.close()

        return occurrences_from_rows(
            dicts, default_lead_days=self._default_lead_days
        )
